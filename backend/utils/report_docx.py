import docx
from docx.shared import Inches, Pt, RGBColor
import os
from backend.config import settings

def generate_docx_report(doc_name: str, summary: str, stats: dict, charts: list[str], output_path: str):
    """
    Generates a DOCX (Microsoft Word) document report containing file summary, stats, and charts.
    """
    doc = docx.Document()
    
    # Document Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("Summarizerr Report: " + doc_name)
    title_run.font.name = 'Arial'
    title_run.font.size = Pt(18)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(79, 70, 229) # Indigo #4f46e5
    
    # Subtitle / Generator stamp
    subtitle_p = doc.add_paragraph()
    subtitle_run = subtitle_p.add_run("AI-Generated Document Summary & Analytical Insights")
    subtitle_run.font.name = 'Arial'
    subtitle_run.font.size = Pt(11)
    subtitle_run.font.italic = True
    subtitle_run.font.color.rgb = RGBColor(128, 128, 128)

    # 1. File Metadata Table
    doc.add_heading("1. File Metadata Information", level=1)
    
    file_type = stats.get("file_type", "").upper()
    text_stats = stats.get("text_stats", {})
    
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Light Shading Accent 1'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Property'
    hdr_cells[1].text = 'Value'
    
    metadata_fields = [
        ("File Name", doc_name),
        ("File Format", file_type),
    ]
    
    if text_stats:
        metadata_fields.extend([
            ("Word Count", f"{text_stats.get('word_count', 0)} words"),
            ("Character Count", f"{text_stats.get('char_count', 0)} characters"),
            ("Estimated Reading Time", f"{text_stats.get('reading_time_minutes', 0)} minutes")
        ])
        
    if stats.get("is_tabular") and "tabular_stats" in stats:
        t_stats = stats["tabular_stats"]
        metadata_fields.extend([
            ("Tabular Dimensions", f"{t_stats.get('row_count', 0)} rows x {t_stats.get('column_count', 0)} columns")
        ])
        
    for prop, val in metadata_fields:
        row_cells = table.add_row().cells
        row_cells[0].text = prop
        row_cells[1].text = val
        
    doc.add_paragraph() # Spacing

    # 2. Summary
    doc.add_heading("2. AI Generated Summary", level=1)
    
    if summary:
        for line in summary.split("\n"):
            line = line.strip()
            if not line:
                continue
                
            if line.startswith("###"):
                heading_text = line.replace("###", "").strip()
                doc.add_heading(heading_text, level=3)
            elif line.startswith("##"):
                heading_text = line.replace("##", "").strip()
                doc.add_heading(heading_text, level=2)
            elif line.startswith("#"):
                heading_text = line.replace("#", "").strip()
                doc.add_heading(heading_text, level=1)
            elif line.startswith("-") or line.startswith("*"):
                bullet_text = line[1:].strip().replace("**", "").replace("*", "")
                doc.add_paragraph(bullet_text, style='List Bullet')
            else:
                normal_text = line.replace("**", "").replace("*", "")
                doc.add_paragraph(normal_text)
    else:
        doc.add_paragraph("No summary available for this document.")

    doc.add_paragraph() # Spacing

    # 3. Visualizations / Charts
    if charts:
        doc.add_page_break()
        doc.add_heading("3. Visualizations & Analytical Charts", level=1)
        doc.add_paragraph("The charts below show data patterns and distributions found within the document:")
        
        charts_dir = os.path.join(settings.UPLOAD_DIR, "charts")
        for filename in charts:
            chart_path = os.path.join(charts_dir, filename)
            if os.path.exists(chart_path):
                doc.add_paragraph(f"Figure: {filename}")
                doc.add_picture(chart_path, width=Inches(5.8))
                doc.add_paragraph()

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
