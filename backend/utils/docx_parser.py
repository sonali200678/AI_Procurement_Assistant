import docx
import os

def parse_file(file_path: str) -> dict:
    """
    Parses a DOCX (Word) file and returns its text content and metadata.
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found at {file_path}")

    text_content = []
    
    try:
        doc = docx.Document(file_path)
        
        # Extract text from paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text_content.append(para.text)

        # Extract text from tables
        table_count = len(doc.tables)
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_content.append(" | ".join(row_text))

        full_text = "\n".join(text_content)
        
        # Prepare metadata
        metadata = {
            "paragraph_count": len(doc.paragraphs),
            "table_count": table_count,
            "word_count": len(full_text.split())
        }

        # Check document core properties if available
        try:
            metadata["title"] = doc.core_properties.title or ""
            metadata["author"] = doc.core_properties.author or ""
        except:
            metadata["title"] = ""
            metadata["author"] = ""

        return {
            "text": full_text,
            "metadata": metadata
        }

    except Exception as e:
        raise ValueError(f"Failed to parse DOCX file: {str(e)}")